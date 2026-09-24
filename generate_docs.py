import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_credentials_document(output_path):
    doc = docx.Document()

    # Set Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("SafePassage AI — System Login Credentials")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(14, 116, 144) # Teal/Cyan

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle_p.add_run("School & College Cab Management Network | Official Access Directory")
    subtitle_run.font.name = "Calibri"
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Overview Callout
    callout_table = doc.add_table(rows=1, cols=1)
    callout_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = callout_table.cell(0, 0)
    set_cell_background(cell, "F0F9FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    cp = cell.paragraphs[0]
    c_run1 = cp.add_run("🔐 STRICT AUTHENTICATION ACTIVE: ")
    c_run1.bold = True
    c_run1.font.color.rgb = RGBColor(3, 105, 161)
    c_run2 = cp.add_run("All applications (Parent App, Driver Cockpit, Cab Owner, Student, Working Pro, and Super Admin) now require entering credentials manually. Use the verified usernames and passwords below.")
    c_run2.font.color.rgb = RGBColor(51, 65, 85)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Heading 1: Master Credentials Table
    h1 = doc.add_paragraph()
    h1_run = h1.add_run("1. Master Credentials Reference Table")
    h1_run.font.name = "Calibri"
    h1_run.font.size = Pt(15)
    h1_run.font.bold = True
    h1_run.font.color.rgb = RGBColor(30, 41, 59)

    # Master Table
    headers = ["Role / Persona", "Username", "Password", "Email / Identifier", "Mobile / Phone"]
    data = [
        ("Parent", "priya_parent", "parent123", "priya.parent@gmail.com", "+91 98401 77889"),
        ("Cab Driver", "kumar_driver", "driver123", "kumar.driver@cabs.com", "+91 98401 23456"),
        ("Cab Owner", "ravi_owner", "owner123", "ravi.owner@chennaicabs.com", "+91 98401 55667"),
        ("Student", "ananya_student", "student123", "ananya.v@college.edu", "+91 98401 88990"),
        ("Corporate Pro", "rohit_pro", "pro123", "rohit.m@infosys.com", "+91 98401 11223"),
        ("Super Admin", "admin", "admin123", "admin@safepassage.ai", "+91 90000 00001")
    ]

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    col_widths = [Inches(1.3), Inches(1.3), Inches(1.1), Inches(1.8), Inches(1.3)]

    # Header Row
    for col_idx, text in enumerate(headers):
        hcell = table.cell(0, col_idx)
        hcell.width = col_widths[col_idx]
        set_cell_background(hcell, "1E293B")
        set_cell_margins(hcell, top=100, bottom=100, left=120, right=120)
        hp = hcell.paragraphs[0]
        hrun = hp.add_run(text)
        hrun.font.bold = True
        hrun.font.size = Pt(10)
        hrun.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, row_data in enumerate(data):
        bg = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            dcell = table.cell(row_idx + 1, col_idx)
            dcell.width = col_widths[col_idx]
            set_cell_background(dcell, bg)
            set_cell_margins(dcell, top=80, bottom=80, left=120, right=120)
            dp = dcell.paragraphs[0]
            drun = dp.add_run(text)
            drun.font.size = Pt(9.5)
            if col_idx == 0:
                drun.font.bold = True
                drun.font.color.rgb = RGBColor(15, 23, 42)
            elif col_idx == 2:
                drun.font.bold = True
                drun.font.color.rgb = RGBColor(2, 132, 199)
            else:
                drun.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Heading 2: Role Details & Scopes
    h2 = doc.add_paragraph()
    h2_run = h2.add_run("2. Role Profiles & Associated Modules")
    h2_run.font.name = "Calibri"
    h2_run.font.size = Pt(15)
    h2_run.font.bold = True
    h2_run.font.color.rgb = RGBColor(30, 41, 59)

    roles_info = [
        {
            "role": "Parent Persona (Priya Sharma)",
            "user": "priya_parent",
            "pass": "parent123",
            "scope": "Live tracking of child rides, SOS alerts, child journey passport, transport booking, monthly fee payments, driver profile verification.",
            "school": "Oakridge International / ABC Matriculation School",
            "color": "0284C7"
        },
        {
            "role": "Cab Driver Persona (Kumar Swamy)",
            "user": "kumar_driver",
            "pass": "driver123",
            "scope": "Trip navigation cockpit, GPS streaming, student boarding attendance checklist, guardian OTP verification, anti-abandonment child sweep checks.",
            "vehicle": "Mercedes / Force Traveller (TN 01 AB 1234)",
            "color": "059669"
        },
        {
            "role": "Cab Owner & Fleet Master (Ravi Networks)",
            "user": "ravi_owner",
            "pass": "owner123",
            "scope": "Fleet vehicle management, vehicle assignment to routes, driver telematics scoring, payout ledger, breakdown backup vehicle dispatch.",
            "fleet": "Chennai IT & School Transport (8 Active Cabs)",
            "color": "D97706"
        },
        {
            "role": "Student Persona (Ananya Verma)",
            "user": "ananya_student",
            "pass": "student123",
            "scope": "Digital commute pass (NFC/QR/Barcode), daily shuttle route map, live ETA, SOS emergency trigger, attendance boarding log.",
            "institution": "SRM Institute of Tech / Loyola College",
            "color": "7C3AED"
        },
        {
            "role": "Corporate Professional (Rohit Menon)",
            "user": "rohit_pro",
            "pass": "pro123",
            "scope": "Morning & evening shift corporate cab booking, recurring ride packages, invoice download, real-time vehicle radar.",
            "company": "Infosys Mahindra World City",
            "color": "0891B2"
        },
        {
            "role": "Super Admin Master (Root Administrator)",
            "user": "admin",
            "pass": "admin123",
            "scope": "Full platform governance, AI route matching orchestration, pending driver/parent KYC approvals, invoice desk, safety incident dispatch.",
            "permission": "ALL_SUPER_ADMIN_PERMISSIONS",
            "color": "DC2626"
        }
    ]

    for item in roles_info:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_after = Pt(2)
        r_title = p_item.add_run(f"• {item['role']}\n")
        r_title.bold = True
        r_title.font.size = Pt(11)
        r_title.font.color.rgb = RGBColor(15, 23, 42)

        r_cred = p_item.add_run(f"   Credentials: Username: {item['user']}  |  Password: {item['pass']}\n")
        r_cred.bold = True
        r_cred.font.size = Pt(10)
        r_cred.font.color.rgb = RGBColor(2, 132, 199)

        r_desc = p_item.add_run(f"   Capabilities: {item['scope']}\n")
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # Heading 3: Emergency OTP Codes
    h3 = doc.add_paragraph()
    h3_run = h3.add_run("3. Emergency 2FA & Offline Bypass OTP Codes")
    h3_run.font.name = "Calibri"
    h3_run.font.size = Pt(15)
    h3_run.font.bold = True
    h3_run.font.color.rgb = RGBColor(30, 41, 59)

    p_otp = doc.add_paragraph()
    p_otp.add_run("• Emergency / Forgot Password Simulated OTP: ").bold = True
    p_otp.add_run("849201\n")
    p_otp.add_run("• Guardian Pickup Verification Default PIN: ").bold = True
    p_otp.add_run("4829\n")
    p_otp.add_run("• Single Trip Quick Booking OTP: ").bold = True
    p_otp.add_run("4-digit dynamic code (45s window)")

    # Footer note
    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_foot = p_foot.add_run("— Generated securely for SafePassage AI Administration —")
    r_foot.font.size = Pt(9)
    r_foot.font.italic = True
    r_foot.font.color.rgb = RGBColor(148, 163, 184)

    # Save
    doc.save(output_path)
    print(f"Document successfully created at {output_path}")

if __name__ == "__main__":
    out1 = r"c:\Users\DELL\OneDrive\Documents\school_college_cab_managements_software\SafePassage_AI_Login_Credentials.docx"
    out2 = r"c:\Users\DELL\OneDrive\Documents\school_college_cab_managements_software\docs\SafePassage_AI_Login_Credentials.docx"
    os.makedirs(r"c:\Users\DELL\OneDrive\Documents\school_college_cab_managements_software\docs", exist_ok=True)
    create_credentials_document(out1)
    create_credentials_document(out2)
